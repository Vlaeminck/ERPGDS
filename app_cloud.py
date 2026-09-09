import os
import sys
import io
import re
from datetime import datetime
from functools import wraps
from flask import Flask, render_template, jsonify, request, session, send_file
from dotenv import load_dotenv

# Cargar variables de entorno si existe .env
load_dotenv()

import db_manager
import firebase_sync

try:
    db_manager.init_db()
except Exception as e:
    print(f"[CloudApp] Error al inicializar base de datos: {e}")

# Iniciar motor de sincronización con la nube (Intervalo: 5 segundos)
try:
    firebase_sync.start_sync_engine(interval=5)
except Exception as e:
    print(f"[CloudApp] Error al iniciar Firebase Sync Engine: {e}")

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'gds-erp-cloud-stakeholder-secret-key-2026')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16 MB max

STAKEHOLDER_PIN = os.environ.get('STAKEHOLDER_PIN', '1234')


# ==========================================
# DECORADOR Y RUTAS DE AUTENTICACIÓN
# ==========================================

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('authenticated'):
            return jsonify({"error": "No autenticado", "require_auth": True}), 401
        return f(*args, **kwargs)
    return decorated_function


@app.route('/')
def index():
    return render_template('stakeholders.html')


@app.route('/api/auth/check', methods=['GET'])
def api_auth_check():
    return jsonify({
        "authenticated": bool(session.get('authenticated', False))
    })


@app.route('/api/auth/login', methods=['POST'])
def api_auth_login():
    data = request.json or {}
    pin = str(data.get('pin', '')).strip()
    
    if pin and pin == STAKEHOLDER_PIN:
        session['authenticated'] = True
        session.permanent = True
        return jsonify({"success": True, "message": "Acceso concedido"})
    return jsonify({"success": False, "message": "PIN incorrecto"}), 401


@app.route('/api/auth/logout', methods=['POST'])
def api_auth_logout():
    session.clear()
    return jsonify({"success": True, "message": "Sesión cerrada"})


# ==========================================
# RUTAS DE ESTADO CLOUD & FIREBASE SYNC
# ==========================================

@app.route('/api/firebase/status', methods=['GET'])
def api_firebase_status():
    return jsonify(firebase_sync.get_sync_status())


@app.route('/api/firebase/sync_now', methods=['POST'])
def api_firebase_sync_now():
    firebase_sync.reconcile_with_firestore()
    firebase_sync.sync_cycle()
    return jsonify(firebase_sync.get_sync_status())


@app.route('/api/firebase/full_resync', methods=['POST', 'GET'])
def api_firebase_full_resync():
    reconciled = firebase_sync.reconcile_with_firestore()
    firebase_sync.pull_remote_changes(force_full=True)
    firebase_sync.push_local_changes()
    return jsonify({"success": True, "reconciled": reconciled, "status": firebase_sync.get_sync_status()})


@app.route('/api/meses_disponibles', methods=['GET'])
def api_meses_disponibles():
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    
    queries = [
        "SELECT DISTINCT strftime('%Y-%m', fecha_emision) FROM arca_compras_csv WHERE fecha_emision IS NOT NULL AND fecha_emision != ''",
        "SELECT DISTINCT strftime('%Y-%m', fecha) FROM proveedores_cuentas_pagar WHERE fecha IS NOT NULL AND fecha != ''"
    ]
    
    meses_set = set()
    for q in queries:
        try:
            cursor.execute(q)
            for r in cursor.fetchall():
                if r[0] and len(r[0]) == 7:
                    meses_set.add(r[0])
        except Exception:
            pass
            
    now = datetime.now()
    cur_month = now.strftime('%Y-%m')
    meses_set.add(cur_month)
    
    for i in range(1, 13):
        meses_set.add(f"{now.year}-{i:02d}")
    
    next_m = (now.month % 12) + 1
    next_y = now.year + (1 if now.month == 12 else 0)
    meses_set.add(f"{next_y}-{next_m:02d}")
    
    conn.close()
    lista_meses = sorted(list(meses_set), reverse=True)
    return jsonify({"meses": lista_meses, "mes_actual": cur_month})


# ==========================================
# RUTAS COMPRAS ARCA
# ==========================================

@app.route('/api/arca_compras', methods=['GET'])
def api_arca_compras():
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    mes = request.args.get('mes')
    
    if mes and mes != 'all':
        cursor.execute("SELECT * FROM arca_compras_csv WHERE mes = ? ORDER BY fecha_emision DESC", (mes,))
        rows = [dict(r) for r in cursor.fetchall()]
        cursor.execute("SELECT SUM(imp_total), SUM(total_iva) FROM arca_compras_csv WHERE mes = ?", (mes,))
    else:
        cursor.execute("SELECT * FROM arca_compras_csv ORDER BY fecha_emision DESC")
        rows = [dict(r) for r in cursor.fetchall()]
        cursor.execute("SELECT SUM(imp_total), SUM(total_iva) FROM arca_compras_csv")
    
    tot_imp, tot_iva = cursor.fetchone()
    tot_imp = tot_imp or 0
    tot_iva = tot_iva or 0
    
    AFIP_NC_CODES = {'3', '8', '13', '15', '53', '03', '08', '003', '008', '013', '053'}
    def _is_nc(tipo):
        t = str(tipo or '').strip().lower()
        return t in AFIP_NC_CODES or ('nota' in t and ('cr' in t or 'credito' in t)) or t == 'nc'

    for r in rows:
        r['is_nc'] = _is_nc(r.get('tipo_comprobante'))

    pendientes = sum(1 for r in rows if r.get('estado') == 'Pendiente')
    pagados_count = sum(1 for r in rows if r.get('estado') == 'Pagado')
    pagados_total = sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado')
    retroactivas_count = sum(1 for r in rows if r.get('es_retroactiva') == 1)
    nc_count = sum(1 for r in rows if r.get('is_nc'))

    # Métricas por método de pago para stakeholders
    pagos_por_metodo = {
        "Efectivo": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and 'Efectivo' in str(r.get('metodo_pago', ''))),
        "Galicia": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and 'Galicia' in str(r.get('metodo_pago', ''))),
        "Mercado Pago": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and 'Mercado' in str(r.get('metodo_pago', ''))),
        "Tarjeta crédito": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and 'Tarjeta' in str(r.get('metodo_pago', '')))
    }

    # Métricas por categoría de pago
    pagos_por_categoria = {
        "Pinamar": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and r.get('categoria_pago') == 'Pinamar'),
        "Leloir": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and r.get('categoria_pago') == 'Leloir'),
        "Socios": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and r.get('categoria_pago') == 'Socios'),
        "Sin Categorizar": sum(r.get('imp_total', 0) for r in rows if r.get('estado') == 'Pagado' and not r.get('categoria_pago'))
    }

    alias_map = db_manager.get_suppliers_alias_map()
    conn.close()

    return jsonify({
        "compras": rows,
        "alias_map": alias_map,
        "resumen": {
            "total_compras": len(rows),
            "total_importe": tot_imp,
            "total_iva": tot_iva,
            "pendientes": pendientes,
            "pagados": pagados_count,
            "pagados_total": pagados_total,
            "retroactivas": retroactivas_count,
            "notas_credito": nc_count,
            "pagos_por_metodo": pagos_por_metodo,
            "pagos_por_categoria": pagos_por_categoria
        }
    })


@app.route('/api/arca_compras/<int:item_id>/marcar_pago', methods=['POST'])
def api_arca_marcar_pago(item_id):
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    data = request.json or {}
    metodo = str(data.get('metodo_pago', '')).strip()
    categoria_pago = str(data.get('categoria_pago', '')).strip()
    fecha_pago = data.get('fecha_pago', datetime.now().strftime('%Y-%m-%d'))
    now_iso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    if not categoria_pago or categoria_pago not in ['Pinamar', 'Leloir', 'Socios']:
        conn.close()
        return jsonify({"success": False, "message": "Debe seleccionar una Categoría de Pago obligatoria (Pinamar, Leloir o Socios)."}), 400

    if not metodo:
        conn.close()
        return jsonify({"success": False, "message": "Debe seleccionar un Método de Pago válido."}), 400
    
    cursor.execute("SELECT * FROM arca_compras_csv WHERE id=?", (item_id,))
    arca_row = cursor.fetchone()
    
    cursor.execute(
        "UPDATE arca_compras_csv SET estado='Pagado', metodo_pago=?, categoria_pago=?, fecha_pago=?, updated_at=?, sync_status=0 WHERE id=?",
        (metodo, categoria_pago, fecha_pago, now_iso, item_id)
    )
    
    if arca_row:
        pv_str = str(arca_row['punto_venta']).strip()
        nro_str = str(arca_row['nro_comprobante']).strip()
        pv = int(pv_str) if pv_str.isdigit() else 0
        nro = int(nro_str) if nro_str.isdigit() else 0
        if pv > 0 and nro > 0:
            cp_metodo = 'Caja Chica'
            if 'Banco' in metodo or 'Galicia' in metodo: cp_metodo = 'Banco'
            if 'Mercado' in metodo: cp_metodo = 'MercadoPago'
            if 'Tarjeta' in metodo: cp_metodo = 'Tarjeta'
            
            cursor.execute("SELECT id, factura_numero FROM proveedores_cuentas_pagar")
            for cp_row in cursor.fetchall():
                cp_fn = str(cp_row['factura_numero'])
                match = re.search(r'(\d+)\s*-\s*(\d+)', cp_fn)
                if match and int(match.group(1)) == pv and int(match.group(2)) == nro:
                    cursor.execute('''
                        UPDATE proveedores_cuentas_pagar
                        SET estado='Pagado', monto_pagado=monto_total, fecha_pago=?, medio_pago=?, categoria_pago=?, updated_at=?, sync_status=0
                        WHERE id=?
                    ''', (fecha_pago, cp_metodo, categoria_pago, now_iso, cp_row['id']))
    conn.commit()
    conn.close()
    
    # Notificar a Firebase
    firebase_sync.sync_cycle()
    return jsonify({"success": True})


@app.route('/api/arca_compras/<int:item_id>/desmarcar_pago', methods=['POST'])
def api_arca_desmarcar_pago(item_id):
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    now_iso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cursor.execute("SELECT * FROM arca_compras_csv WHERE id=?", (item_id,))
    arca_row = cursor.fetchone()

    cursor.execute(
        "UPDATE arca_compras_csv SET estado='Pendiente', metodo_pago='', fecha_pago='', updated_at=?, sync_status=0 WHERE id=?",
        (now_iso, item_id)
    )
    
    if arca_row:
        pv_str = str(arca_row['punto_venta']).strip()
        nro_str = str(arca_row['nro_comprobante']).strip()
        pv = int(pv_str) if pv_str.isdigit() else 0
        nro = int(nro_str) if nro_str.isdigit() else 0
        if pv > 0 and nro > 0:
            cursor.execute("SELECT id, factura_numero FROM proveedores_cuentas_pagar")
            for cp_row in cursor.fetchall():
                cp_fn = str(cp_row['factura_numero'])
                match = re.search(r'(\d+)\s*-\s*(\d+)', cp_fn)
                if match and int(match.group(1)) == pv and int(match.group(2)) == nro:
                    cursor.execute('''
                        UPDATE proveedores_cuentas_pagar
                        SET estado='Pendiente', monto_pagado=0, fecha_pago='', medio_pago='', updated_at=?, sync_status=0
                        WHERE id=?
                    ''', (now_iso, cp_row['id']))
                    
    conn.commit()
    conn.close()
    firebase_sync.sync_cycle()
    return jsonify({"success": True})


@app.route('/api/arca_compras/<int:item_id>/marcar_recibida', methods=['POST'])
def api_arca_marcar_recibida(item_id):
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    now_iso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cursor.execute(
        "UPDATE arca_compras_csv SET factura_recibida=1, updated_at=?, sync_status=0 WHERE id=?",
        (now_iso, item_id)
    )
    conn.commit()
    conn.close()
    firebase_sync.sync_cycle()
    return jsonify({"success": True})


@app.route('/api/test/marcar_mes_pagado', methods=['POST'])
def api_test_marcar_mes_pagado():
    data = request.json or {}
    mes = data.get('mes')
    categoria_pago = str(data.get('categoria_pago', '')).strip()
    metodo_pago = str(data.get('metodo_pago', 'Galicia')).strip() or 'Galicia'
    
    if not mes:
        return jsonify({"success": False, "message": "Mes no proporcionado"}), 400
    
    if not categoria_pago or categoria_pago not in ['Pinamar', 'Leloir', 'Socios']:
        return jsonify({"success": False, "message": "Debe seleccionar una Categoría de Pago obligatoria (Pinamar, Leloir o Socios)."}), 400
        
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    now_iso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cursor.execute('''
        UPDATE arca_compras_csv 
        SET estado = 'Pagado', metodo_pago = ?, categoria_pago = ?, fecha_pago = date('now'), updated_at=?, sync_status=0
        WHERE mes = ? AND estado != 'Pagado'
    ''', (metodo_pago, categoria_pago, now_iso, mes))
    
    actualizados = cursor.rowcount
    conn.commit()
    conn.close()
    firebase_sync.sync_cycle()
    return jsonify({
        "success": True, 
        "actualizados": actualizados,
        "message": f"Se marcaron como pagadas {actualizados} compras de {mes} ({categoria_pago}) con {metodo_pago}."
    })


# ==========================================
# RUTAS EXPORTACIÓN E IMPORTACIÓN EXCEL
# ==========================================

@app.route('/api/arca_compras/export_excel', methods=['GET'])
def api_arca_export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    conn = db_manager.get_connection()
    cursor = conn.cursor()
    try:
        mes = request.args.get('mes')
        if mes and mes != 'all':
            cursor.execute("SELECT * FROM arca_compras_csv WHERE mes = ? ORDER BY fecha_emision DESC", (mes,))
            filename_part = mes
        else:
            cursor.execute("SELECT * FROM arca_compras_csv ORDER BY fecha_emision DESC")
            filename_part = "todas"
        
        rows = [dict(r) for r in cursor.fetchall()]
    finally:
        conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Compras ARCA ({filename_part})"[:31]

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data_font = Font(name='Segoe UI', size=10)
    bold_font = Font(name='Segoe UI', size=10, bold=True)
    
    fill_pagado = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
    font_pagado = Font(name='Segoe UI', size=10, bold=True, color='166534')
    fill_pendiente = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    font_pendiente = Font(name='Segoe UI', size=10, bold=True, color='92400E')

    thin_border_side = Side(style='thin', color='CBD5E1')
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    headers = [
        ("ID", 8),
        ("Fecha Emisión", 14),
        ("Denominación Emisor", 38),
        ("CUIT Emisor", 16),
        ("Punto Venta", 13),
        ("Nro. Comprobante", 17),
        ("Tipo Comprobante", 17),
        ("CAE", 18),
        ("Total IVA ($)", 15),
        ("Imp. Total ($)", 16),
        ("Factura Recibida", 16),
        ("Estado", 14),
        ("Método de Pago", 20),
        ("Fecha Pago", 14)
    ]

    for col_idx, (h_name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 28

    for row_idx, r in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 20
        estado_str = str(r.get('estado') or 'Pendiente').strip()
        recibida_str = "Sí" if r.get('factura_recibida') else "No"

        row_values = [
            r.get('id'),
            r.get('fecha_emision') or '',
            r.get('denominacion_emisor') or '',
            str(r.get('nro_doc_emisor') or ''),
            str(r.get('punto_venta') or ''),
            str(r.get('nro_comprobante') or ''),
            r.get('tipo_comprobante') or '',
            str(r.get('cae') or ''),
            float(r.get('total_iva') or 0),
            float(r.get('imp_total') or 0),
            recibida_str,
            estado_str,
            r.get('metodo_pago') or '',
            r.get('fecha_pago') or ''
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = cell_border

            if col_idx == 1:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.font = bold_font
            elif col_idx in (4, 5, 6, 7, 8):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in (9, 10):
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '$#,##0.00'
                cell.font = bold_font
            elif col_idx == 11:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if val == "Sí":
                    cell.font = Font(name='Segoe UI', size=10, bold=True, color='047857')
            elif col_idx == 12:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if estado_str == 'Pagado':
                    cell.fill = fill_pagado
                    cell.font = font_pagado
                else:
                    cell.fill = fill_pendiente
                    cell.font = font_pendiente
            elif col_idx in (13, 14):
                cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"Compras_ARCA_{filename_part}_{timestamp}.xlsx"
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


@app.route('/api/arca_compras/import_excel', methods=['POST'])
def api_arca_import_excel():
    import openpyxl

    if 'file' not in request.files:
        return jsonify({"success": False, "message": "No se seleccionó ningún archivo"}), 400
        
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({"success": False, "message": "Archivo no válido"}), 400

    filename_lower = file.filename.lower()
    if not (filename_lower.endswith('.xlsx') or filename_lower.endswith('.xls')):
        return jsonify({"success": False, "message": "El formato del archivo debe ser Excel (.xlsx)"}), 400

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active
    except Exception as e:
        return jsonify({"success": False, "message": f"Error al abrir el archivo Excel: {str(e)}"}), 400

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return jsonify({"success": False, "message": "El archivo Excel está vacío o no contiene filas de datos."}), 400

    raw_header = [str(col or '').strip().lower() for col in rows[0]]

    def find_col_idx(patterns):
        for pat in patterns:
            for idx, h in enumerate(raw_header):
                if pat in h:
                    return idx
        return -1

    idx_id = find_col_idx(['id', 'código', 'codigo'])
    idx_fecha = find_col_idx(['fecha emisión', 'fecha emision', 'fecha_emision', 'fecha'])
    idx_denom = find_col_idx(['denominación', 'denominacion', 'emisor', 'proveedor', 'razon social', 'razón social'])
    idx_cuit = find_col_idx(['cuit', 'nro doc', 'nro_doc', 'documento'])
    idx_pv = find_col_idx(['punto venta', 'punto_venta', 'pto venta', 'pto_venta', 'pv'])
    idx_nro_comp = find_col_idx(['nro. comprobante', 'nro comprobante', 'nro_comprobante', 'comprobante', 'factura', 'numero'])
    idx_tipo = find_col_idx(['tipo comprobante', 'tipo_comprobante', 'tipo'])
    idx_cae = find_col_idx(['cae', 'cód. autorización', 'cod autorizacion'])
    idx_iva = find_col_idx(['iva', 'total iva'])
    idx_total = find_col_idx(['imp. total', 'imp total', 'importe total', 'total'])
    idx_recibida = find_col_idx(['factura recibida', 'recibida', 'recibido'])
    idx_estado = find_col_idx(['estado', 'estado pago', 'estado_pago'])
    idx_metodo = find_col_idx(['método de pago', 'metodo de pago', 'metodo_pago', 'medio de pago', 'medio_pago', 'forma de pago'])
    idx_fecha_pago = find_col_idx(['fecha pago', 'fecha_pago', 'fecha de pago'])

    conn = db_manager.get_connection()
    cursor = conn.cursor()
    now_iso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    today_iso = datetime.now().strftime('%Y-%m-%d')

    actualizados = 0
    nuevos = 0
    filas_ignoradas = 0

    try:
        cursor.execute("SELECT id, factura_numero, monto_total FROM proveedores_cuentas_pagar")
        cp_rows = cursor.fetchall()
        cp_map = []
        for cp in cp_rows:
            cp_fn = str(cp['factura_numero'])
            match = re.search(r'(\d+)\s*-\s*(\d+)', cp_fn)
            if match:
                cp_map.append({
                    'id': cp['id'],
                    'pv': int(match.group(1)),
                    'nro': int(match.group(2)),
                    'monto': cp['monto_total']
                })

        for row in rows[1:]:
            if not any(row):
                continue

            def get_cell(idx, default=''):
                if idx < 0 or idx >= len(row) or row[idx] is None:
                    return default
                return row[idx]

            raw_id = get_cell(idx_id)
            row_id = None
            if raw_id is not None:
                try:
                    row_id = int(raw_id)
                except (ValueError, TypeError):
                    row_id = None

            raw_fecha_em = get_cell(idx_fecha)
            fecha_em_str = ''
            if isinstance(raw_fecha_em, datetime):
                fecha_em_str = raw_fecha_em.strftime('%Y-%m-%d')
            elif hasattr(raw_fecha_em, 'strftime'):
                fecha_em_str = raw_fecha_em.strftime('%Y-%m-%d')
            elif raw_fecha_em:
                raw_s = str(raw_fecha_em).strip()
                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
                    try:
                        d = datetime.strptime(raw_s, fmt)
                        fecha_em_str = d.strftime('%Y-%m-%d')
                        break
                    except Exception:
                        pass
                if not fecha_em_str:
                    fecha_em_str = raw_s

            denom = str(get_cell(idx_denom, '')).strip()
            cuit = str(get_cell(idx_cuit, '')).strip()
            pv_val = str(get_cell(idx_pv, '')).strip()
            nro_comp_val = str(get_cell(idx_nro_comp, '')).strip()
            tipo_val = str(get_cell(idx_tipo, '')).strip()
            cae_val = str(get_cell(idx_cae, '')).strip()

            try:
                raw_iva = get_cell(idx_iva, 0)
                if isinstance(raw_iva, str):
                    raw_iva = raw_iva.replace('$', '').replace('.', '').replace(',', '.').strip()
                iva_val = float(raw_iva or 0)
            except Exception:
                iva_val = 0.0

            try:
                raw_tot = get_cell(idx_total, 0)
                if isinstance(raw_tot, str):
                    raw_tot = raw_tot.replace('$', '').replace('.', '').replace(',', '.').strip()
                tot_val = float(raw_tot or 0)
            except Exception:
                tot_val = 0.0

            raw_estado = str(get_cell(idx_estado, '')).strip().lower()
            if raw_estado in ('pagado', 'pagada', 'pago', 'si', 'sí', '1', 'true', 'ok', 's'):
                estado_norm = 'Pagado'
            elif raw_estado in ('pendiente', 'no', '0', 'false', 'debe', 'n', ''):
                estado_norm = 'Pendiente'
            else:
                estado_norm = 'Pagado' if 'pag' in raw_estado else 'Pendiente'

            raw_metodo = str(get_cell(idx_metodo, '')).strip()
            metodo_norm = ''
            if estado_norm == 'Pagado':
                met_lower = raw_metodo.lower()
                if any(x in met_lower for x in ('efect', 'caja', 'cash', 'm1')):
                    metodo_norm = 'Efectivo'
                elif any(x in met_lower for x in ('gali', 'banc', 'transf', 'cuenta')):
                    metodo_norm = 'Galicia'
                elif any(x in met_lower for x in ('mercado', 'mp', 'm2', 'digital')):
                    metodo_norm = 'Mercado Pago'
                elif any(x in met_lower for x in ('tarj', 'tc', 'cred', 'deb', 'visa', 'master')):
                    metodo_norm = 'Tarjeta crédito'
                elif raw_metodo:
                    metodo_norm = raw_metodo
                else:
                    metodo_norm = 'Efectivo'

            raw_fecha_pago = get_cell(idx_fecha_pago)
            fecha_pago_norm = ''
            if estado_norm == 'Pagado':
                if isinstance(raw_fecha_pago, datetime):
                    fecha_pago_norm = raw_fecha_pago.strftime('%Y-%m-%d')
                elif hasattr(raw_fecha_pago, 'strftime'):
                    fecha_pago_norm = raw_fecha_pago.strftime('%Y-%m-%d')
                elif raw_fecha_pago:
                    raw_p_s = str(raw_fecha_pago).strip()
                    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
                        try:
                            d = datetime.strptime(raw_p_s, fmt)
                            fecha_pago_norm = d.strftime('%Y-%m-%d')
                            break
                        except Exception:
                            pass
                    if not fecha_pago_norm:
                        fecha_pago_norm = raw_p_s
                if not fecha_pago_norm:
                    fecha_pago_norm = today_iso

            raw_recibida = str(get_cell(idx_recibida, '')).strip().lower()
            if raw_recibida in ('sí', 'si', '1', 'true', 'recibida', 'recibido', 's', 'yes', 'y'):
                recibida_norm = 1
            elif raw_recibida in ('no', '0', 'false', 'pendiente', 'n', ''):
                recibida_norm = 0
            else:
                recibida_norm = 1 if 'recib' in raw_recibida or 'si' in raw_recibida else 0

            existing_row = None
            if row_id:
                cursor.execute("SELECT * FROM arca_compras_csv WHERE id = ?", (row_id,))
                existing_row = cursor.fetchone()

            if not existing_row and cae_val:
                cursor.execute("SELECT * FROM arca_compras_csv WHERE cae = ? AND cae != ''", (cae_val,))
                existing_row = cursor.fetchone()

            if not existing_row and cuit and pv_val and nro_comp_val:
                cursor.execute(
                    "SELECT * FROM arca_compras_csv WHERE nro_doc_emisor = ? AND punto_venta = ? AND nro_comprobante = ?",
                    (cuit, pv_val, nro_comp_val)
                )
                existing_row = cursor.fetchone()

            if not existing_row and fecha_em_str and denom and nro_comp_val:
                cursor.execute(
                    "SELECT * FROM arca_compras_csv WHERE fecha_emision = ? AND denominacion_emisor = ? AND nro_comprobante = ?",
                    (fecha_em_str, denom, nro_comp_val)
                )
                existing_row = cursor.fetchone()

            if existing_row:
                target_id = existing_row['id']
                cursor.execute('''
                    UPDATE arca_compras_csv
                    SET estado = ?, metodo_pago = ?, fecha_pago = ?, factura_recibida = ?, updated_at = ?, sync_status = 0
                    WHERE id = ?
                ''', (estado_norm, metodo_norm, fecha_pago_norm, recibida_norm, now_iso, target_id))
                actualizados += 1

                pv_curr = str(existing_row['punto_venta'] or pv_val).strip()
                nro_curr = str(existing_row['nro_comprobante'] or nro_comp_val).strip()
                if pv_curr.isdigit() and nro_curr.isdigit():
                    pv_i = int(pv_curr)
                    nro_i = int(nro_curr)
                    for cp in cp_map:
                        if cp['pv'] == pv_i and cp['nro'] == nro_i:
                            if estado_norm == 'Pagado':
                                cp_m = 'Caja Chica'
                                if 'Galicia' in metodo_norm or 'Banco' in metodo_norm:
                                    cp_m = 'Banco'
                                elif 'Mercado' in metodo_norm:
                                    cp_m = 'MercadoPago'
                                cursor.execute('''
                                    UPDATE proveedores_cuentas_pagar
                                    SET estado='Pagado', monto_pagado=monto_total, fecha_pago=?, medio_pago=?, updated_at=?, sync_status=0
                                    WHERE id=?
                                ''', (fecha_pago_norm, cp_m, now_iso, cp['id']))
                            else:
                                cursor.execute('''
                                    UPDATE proveedores_cuentas_pagar
                                    SET estado='Pendiente', monto_pagado=0, fecha_pago='', medio_pago='', updated_at=?, sync_status=0
                                    WHERE id=?
                                ''', (now_iso, cp['id']))
            else:
                if fecha_em_str and (denom or cuit):
                    mes_calc = fecha_em_str[:7] if len(fecha_em_str) >= 7 else ''
                    cursor.execute('''
                        INSERT INTO arca_compras_csv (
                            fecha_emision, punto_venta, nro_doc_emisor, denominacion_emisor,
                            total_iva, imp_total, mes, estado, factura_recibida,
                            metodo_pago, fecha_pago, cae, nro_comprobante, tipo_comprobante,
                            fecha_importacion
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        fecha_em_str, pv_val, cuit, denom,
                        iva_val, tot_val, mes_calc, estado_norm, recibida_norm,
                        metodo_norm, fecha_pago_norm, cae_val, nro_comp_val, tipo_val,
                        now_iso
                    ))
                    nuevos += 1
                else:
                    filas_ignoradas += 1

        conn.commit()
    finally:
        conn.close()

    firebase_sync.sync_cycle()

    msg_parts = []
    if actualizados > 0:
        msg_parts.append(f"{actualizados} comprobantes actualizados")
    if nuevos > 0:
        msg_parts.append(f"{nuevos} comprobantes nuevos importados")
    if not msg_parts:
        msg_parts.append("No se registraron cambios")

    return jsonify({
        "success": True,
        "actualizados": actualizados,
        "nuevos": nuevos,
        "filas_ignoradas": filas_ignoradas,
        "message": ", ".join(msg_parts) + "."
    })


# ==========================================
# RUTAS CUENTAS POR PAGAR & PROVEEDORES
# ==========================================

@app.route('/api/cuentas_por_pagar', methods=['GET'])
def api_cuentas_por_pagar():
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    mes = request.args.get('mes')
    
    if mes and mes != 'all':
        cursor.execute("SELECT * FROM proveedores_cuentas_pagar WHERE fecha LIKE ? ORDER BY id DESC", (f"{mes}%",))
        rows = [dict(r) for r in cursor.fetchall()]
        
        cursor.execute("SELECT SUM(imp_total) FROM arca_compras_csv WHERE mes = ?", (mes,))
        tot_fact = cursor.fetchone()[0] or 0
        cursor.execute("SELECT SUM(imp_total) FROM arca_compras_csv WHERE mes = ? AND estado = 'Pagado'", (mes,))
        tot_pag = cursor.fetchone()[0] or 0
    else:
        cursor.execute("SELECT * FROM proveedores_cuentas_pagar ORDER BY id DESC")
        rows = [dict(r) for r in cursor.fetchall()]
        
        cursor.execute("SELECT SUM(imp_total) FROM arca_compras_csv")
        tot_fact = cursor.fetchone()[0] or 0
        cursor.execute("SELECT SUM(imp_total) FROM arca_compras_csv WHERE estado = 'Pagado'")
        tot_pag = cursor.fetchone()[0] or 0
        
    pendiente = tot_fact - tot_pag
    conn.close()
    return jsonify({
        "cuentas": rows,
        "resumen": {
            "total_facturado": tot_fact,
            "total_pagado": tot_pag,
            "total_pendiente": pendiente
        }
    })


@app.route('/api/cuentas_por_pagar/registrar_pago', methods=['POST'])
def api_registrar_pago_proveedor():
    conn = db_manager.get_connection()
    cursor = conn.cursor()
    data = request.json or {}
    cuenta_id = data.get('id')
    medio_pago = str(data.get('medio_pago', 'Galicia')).strip() or 'Galicia'
    categoria_pago = str(data.get('categoria_pago', '')).strip()
    fecha_pago = data.get('fecha_pago', datetime.now().strftime('%Y-%m-%d'))
    now_iso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    if not cuenta_id:
        conn.close()
        return jsonify({"error": "Datos inválidos"}), 400

    if not categoria_pago or categoria_pago not in ['Pinamar', 'Leloir', 'Socios']:
        conn.close()
        return jsonify({"error": "Debe seleccionar una Categoría de Pago obligatoria (Pinamar, Leloir o Socios)."}), 400
        
    cursor.execute("SELECT * FROM proveedores_cuentas_pagar WHERE id = ?", (cuenta_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Cuenta no encontrada"}), 404
        
    monto_pago = row['monto_total']
    
    cursor.execute('''
        UPDATE proveedores_cuentas_pagar
        SET monto_pagado = ?, estado = 'Pagado', fecha_pago = ?, medio_pago = ?, categoria_pago = ?, updated_at = ?, sync_status = 0
        WHERE id = ?
    ''', (monto_pago, fecha_pago, medio_pago, categoria_pago, now_iso, cuenta_id))
    
    # Sincronizar pago con ARCA
    match = re.search(r'(\d+)\s*-\s*(\d+)', str(row['factura_numero']))
    if match:
        pv = int(match.group(1))
        nro = int(match.group(2))
        arca_metodo = 'Galicia'
        if 'Efectivo' in medio_pago or 'Caja' in medio_pago: arca_metodo = 'Efectivo'
        elif 'Mercado' in medio_pago: arca_metodo = 'Mercado Pago'
        elif 'Tarjeta' in medio_pago: arca_metodo = 'Tarjeta crédito'

        cursor.execute('''
            UPDATE arca_compras_csv 
            SET estado='Pagado', metodo_pago=?, categoria_pago=?, fecha_pago=?, updated_at=?, sync_status=0
            WHERE CAST(punto_venta AS INTEGER) = ? AND CAST(nro_comprobante AS INTEGER) = ?
        ''', (arca_metodo, categoria_pago, fecha_pago, now_iso, pv, nro))

    conn.commit()
    conn.close()
    firebase_sync.sync_cycle()
    return jsonify({"success": True})


@app.route('/api/suppliers', methods=['GET'])
def api_suppliers():
    rows = db_manager.get_all_unique_suppliers()
    return jsonify({"suppliers": rows})


@app.route('/api/proveedores/alias', methods=['GET', 'POST'])
def api_proveedores_alias():
    if request.method == 'POST':
        data = request.json or {}
        nombre = str(data.get('nombre', '')).strip()
        alias = str(data.get('alias', '')).strip()
        
        if not nombre:
            return jsonify({"success": False, "message": "El nombre del proveedor es obligatorio"}), 400
            
        db_manager.update_supplier_alias(nombre, alias)
        firebase_sync.sync_cycle()
        return jsonify({"success": True, "message": f"Alias guardado correctamente para {nombre}"})
    else:
        rows = db_manager.get_all_unique_suppliers()
        return jsonify({"proveedores": rows})


@app.route('/api/dashboard/stats', methods=['GET'])
def api_dashboard_stats():
    mes = request.args.get('mes')
    conn = db_manager.get_connection()
    cursor = conn.cursor()

    # 1. Evolución mensual (Todos los meses con compras)
    cursor.execute('''
        SELECT mes, 
               SUM(imp_total) as total_facturado,
               SUM(CASE WHEN estado = 'Pagado' THEN imp_total ELSE 0 END) as total_pagado,
               SUM(CASE WHEN estado != 'Pagado' THEN imp_total ELSE 0 END) as total_pendiente
        FROM arca_compras_csv
        WHERE mes IS NOT NULL AND mes != ''
        GROUP BY mes
        ORDER BY mes ASC
    ''')
    evolucion_rows = cursor.fetchall()
    evolucion_mensual = {
        "meses": [r['mes'] for r in evolucion_rows],
        "facturado": [round(float(r['total_facturado'] or 0), 2) for r in evolucion_rows],
        "pagado": [round(float(r['total_pagado'] or 0), 2) for r in evolucion_rows],
        "pendiente": [round(float(r['total_pendiente'] or 0), 2) for r in evolucion_rows]
    }

    # Filtro base para las siguientes consultas
    where_mes = "WHERE mes = ?" if (mes and mes != 'all') else "WHERE 1=1"
    params_mes = (mes,) if (mes and mes != 'all') else ()

    # 2. Desglose por Categoría de Pago
    cursor.execute(f'''
        SELECT COALESCE(NULLIF(categoria_pago, ''), 'Sin Categorizar') as cat,
               SUM(imp_total) as total
        FROM arca_compras_csv
        {where_mes} AND estado = 'Pagado'
        GROUP BY cat
    ''', params_mes)
    cat_rows = cursor.fetchall()
    gastos_categoria = {r['cat']: round(float(r['total'] or 0), 2) for r in cat_rows}

    # 3. Desglose por Método de Pago
    cursor.execute(f'''
        SELECT COALESCE(NULLIF(metodo_pago, ''), 'Sin Definir') as metodo,
               SUM(imp_total) as total
        FROM arca_compras_csv
        {where_mes} AND estado = 'Pagado'
        GROUP BY metodo
    ''', params_mes)
    metodo_rows = cursor.fetchall()
    gastos_metodo = {r['metodo']: round(float(r['total'] or 0), 2) for r in metodo_rows}

    # 4. Top 10 Proveedores con Mayor Gasto
    cursor.execute(f'''
        SELECT denominacion_emisor, SUM(imp_total) as total
        FROM arca_compras_csv
        {where_mes}
        GROUP BY denominacion_emisor
        ORDER BY total DESC
        LIMIT 10
    ''', params_mes)
    top_rows = cursor.fetchall()

    alias_map = db_manager.get_suppliers_alias_map()
    top_proveedores = []
    for r in top_rows:
        razon = r['denominacion_emisor'] or 'Desconocido'
        alias = alias_map.get(razon, '')
        top_proveedores.append({
            "razon_social": razon,
            "display_name": alias if alias else (razon[:30] + '...' if len(razon) > 30 else razon),
            "total": round(float(r['total'] or 0), 2)
        })

    conn.close()

    return jsonify({
        "evolucion_mensual": evolucion_mensual,
        "gastos_categoria": gastos_categoria,
        "gastos_metodo": gastos_metodo,
        "top_proveedores": top_proveedores,
        "mes_seleccionado": mes or 'all'
    })


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)

